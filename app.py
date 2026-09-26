import os
import io
import time
import hashlib
import uuid
from pathlib import Path

import chromadb
import streamlit as st
from dotenv import load_dotenv
from sentence_transformers import SentenceTransformer
from google import genai
from PyPDF2 import PdfReader
from docx import Document

# =========================================================
# CONFIG
# =========================================================
BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
PRIMARY_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
FALLBACK_MODEL = os.getenv("GEMINI_FALLBACK_MODEL", "gemini-3.5-flash-lite")

st.set_page_config(
    page_title="CareerAI",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)

# =========================================================
# UI - high contrast / no white input problem
# =========================================================
st.markdown("""
<style>
#MainMenu, footer, header {visibility:hidden;}

.stApp {
    background:
        radial-gradient(circle at 8% 0%, rgba(37,99,235,.18), transparent 28%),
        radial-gradient(circle at 92% 0%, rgba(139,92,246,.16), transparent 28%),
        #050914;
    color:#f8fafc;
}

.block-container {
    max-width:1450px;
    padding-top:1.2rem;
    padding-bottom:4rem;
}

.hero {
    padding:32px 38px;
    border-radius:24px;
    margin-bottom:22px;
    background:linear-gradient(135deg,#101c35,#24184a);
    border:1px solid rgba(96,165,250,.38);
    box-shadow:0 18px 55px rgba(0,0,0,.30);
}

.main-title {
    font-size:52px;
    font-weight:900;
    line-height:1;
    background:linear-gradient(90deg,#60a5fa,#c084fc,#22d3ee,#34d399);
    -webkit-background-clip:text;
    -webkit-text-fill-color:transparent;
}

.subtitle {font-size:20px;color:#ddd6fe;margin:10px 0;font-weight:750;}
.section-title {font-size:25px;font-weight:850;color:#fff;margin:24px 0 14px;}

.card {
    padding:20px;
    border-radius:17px;
    background:linear-gradient(145deg,#101827,#18253b);
    border:1px solid rgba(148,163,184,.25);
    box-shadow:0 10px 30px rgba(0,0,0,.18);
}

.card h3 {color:#fff;margin-top:0;}
.small-note {color:#cbd5e1;font-size:13px;line-height:1.55;}

.score-card {
    padding:24px;
    border-radius:20px;
    background:linear-gradient(135deg,#111d36,#30245f);
    border:1px solid rgba(167,139,250,.42);
    text-align:center;
}

.big-score {font-size:52px;font-weight:950;color:#67e8f9;}

.score-bar {
    height:12px;
    background:#26344c;
    border-radius:99px;
    overflow:hidden;
    margin:12px 0;
}

.score-fill {
    height:100%;
    background:linear-gradient(90deg,#22d3ee,#60a5fa,#a78bfa,#34d399);
    border-radius:99px;
}

[data-testid="stSidebar"] {
    background:linear-gradient(180deg,#070d1a,#101a2d) !important;
}

[data-testid="stSidebar"] * {color:#f1f5f9 !important;}

.stButton > button {
    border-radius:12px !important;
    min-height:44px !important;
    font-weight:750 !important;
    color:#ffffff !important;
    background:linear-gradient(135deg,#182640,#28385c) !important;
    border:1px solid #52627f !important;
}

.stButton > button:hover {
    border-color:#60a5fa !important;
    background:#20365b !important;
}

/* ---------- File uploader: dark instead of white ---------- */
[data-testid="stFileUploaderDropzone"] {
    background:#101827 !important;
    border:1px dashed #64748b !important;
    border-radius:14px !important;
}

[data-testid="stFileUploaderDropzone"] > div {
    background:#101827 !important;
}

[data-testid="stFileUploaderDropzone"] * {
    color:#f8fafc !important;
    opacity:1 !important;
}

[data-testid="stFileUploaderDropzone"] button {
    background:#243451 !important;
    color:#ffffff !important;
    border:1px solid #64748b !important;
    border-radius:10px !important;
}

[data-testid="stFileUploaderFile"] {
    background:#101827 !important;
    border:1px solid #334155 !important;
}

[data-testid="stFileUploaderFile"] * {
    color:#f8fafc !important;
    opacity:1 !important;
}

/* ---------- Normal text inputs: dark, not white ---------- */
div[data-testid="stTextInput"] input,
div[data-testid="stTextArea"] textarea {
    background:#0f172a !important;
    color:#ffffff !important;
    -webkit-text-fill-color:#ffffff !important;
    caret-color:#67e8f9 !important;
    border:1px solid #475569 !important;
    border-radius:12px !important;
    opacity:1 !important;
}

div[data-testid="stTextInput"] input:focus,
div[data-testid="stTextArea"] textarea:focus {
    background:#111c31 !important;
    color:#ffffff !important;
    -webkit-text-fill-color:#ffffff !important;
    border-color:#60a5fa !important;
    box-shadow:0 0 0 2px rgba(96,165,250,.16) !important;
}

div[data-testid="stTextInput"] input::placeholder,
div[data-testid="stTextArea"] textarea::placeholder {
    color:#94a3b8 !important;
    -webkit-text-fill-color:#94a3b8 !important;
    opacity:1 !important;
}

/* ---------- Chat input: prevent light/disabled-looking box ---------- */
div[data-testid="stChatInput"] {
    background:#0b1220 !important;
    border:1px solid #475569 !important;
    border-radius:16px !important;
    box-shadow:0 8px 25px rgba(0,0,0,.25) !important;
    opacity:1 !important;
}

div[data-testid="stChatInput"] > div {
    background:#0b1220 !important;
    border-radius:16px !important;
    opacity:1 !important;
}

div[data-testid="stChatInput"] textarea {
    background:#0b1220 !important;
    color:#ffffff !important;
    -webkit-text-fill-color:#ffffff !important;
    caret-color:#67e8f9 !important;
    opacity:1 !important;
}

div[data-testid="stChatInput"] textarea::placeholder {
    color:#94a3b8 !important;
    -webkit-text-fill-color:#94a3b8 !important;
    opacity:1 !important;
}

div[data-testid="stChatInput"] button {
    color:#ffffff !important;
    opacity:1 !important;
}

div[data-testid="stChatMessage"],
div[data-testid="stChatMessageContent"] {
    opacity:1 !important;
}

div[data-testid="stChatMessageContent"] * {
    color:#f8fafc !important;
    opacity:1 !important;
}

/* Markdown/result readability */
[data-testid="stMarkdownContainer"] p,
[data-testid="stMarkdownContainer"] li,
[data-testid="stMarkdownContainer"] strong,
[data-testid="stMarkdownContainer"] h1,
[data-testid="stMarkdownContainer"] h2,
[data-testid="stMarkdownContainer"] h3 {
    color:#f8fafc;
}

div[data-testid="stAlert"] {opacity:1 !important;}
</style>
""", unsafe_allow_html=True)

# =========================================================
# SESSION STATE
# =========================================================
def init_state():
    defaults = {
        "user_name": None,
        "messages": [],
        "resume_text": None,
        "resume_name": None,
        "resume_hash": None,
        "jd_text": None,
        "jd_name": None,
        "jd_hash": None,
        "analysis": None,
        "db_ready": False,
        "chroma_client": None,
        "chroma_collection": None,
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value

init_state()

# =========================================================
# NAME-ONLY START
# =========================================================
def show_login():
    st.markdown("""
    <div class="hero">
        <div class="main-title">🎯 CareerAI</div>
        <div class="subtitle">RAG-Powered Resume & Job Description Intelligence</div>
        <div style="color:#f1f5f9;font-size:17px;line-height:1.7;">
            Enter your name to start your career analysis.
        </div>
    </div>
    """, unsafe_allow_html=True)

    _, center, _ = st.columns([1, 1.3, 1])
    with center:
        name = st.text_input(
            "Your name",
            placeholder="e.g. Yuvraj Singh",
            max_chars=100,
        )
        if st.button("🚀 Start CareerAI", use_container_width=True, type="primary"):
            if len(name.strip()) < 2:
                st.warning("Please enter your name.")
            else:
                st.session_state.user_name = name.strip()
                st.rerun()

    st.stop()

if not st.session_state.user_name:
    show_login()

# =========================================================
# MODELS
# =========================================================
@st.cache_resource
def get_embedder():
    return SentenceTransformer("all-MiniLM-L6-v2")

@st.cache_resource
def get_gemini():
    if not GEMINI_API_KEY:
        raise RuntimeError("GEMINI_API_KEY is missing from your .env file.")
    return genai.Client(api_key=GEMINI_API_KEY)

# =========================================================
# SESSION-ISOLATED CHROMA
# =========================================================
def get_chroma():
    """
    Creates an in-memory Chroma client/collection for this
    Streamlit session. Different users do not share documents.
    """
    if st.session_state.chroma_client is None:
        st.session_state.chroma_client = chromadb.Client()

    if st.session_state.chroma_collection is None:
        name = f"careerai_{uuid.uuid4().hex[:12]}"
        st.session_state.chroma_collection = (
            st.session_state.chroma_client.get_or_create_collection(
                name=name,
                metadata={"hnsw:space": "cosine"},
            )
        )

    return st.session_state.chroma_collection

# =========================================================
# FILE EXTRACTION
# =========================================================
def extract_text(filename, file_bytes):
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(file_bytes))
        parts = [(page.extract_text() or "").strip() for page in reader.pages]
        text = "\n\n".join(p for p in parts if p)

    elif suffix == ".docx":
        doc = Document(io.BytesIO(file_bytes))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))

        text = "\n".join(parts)

    elif suffix in (".txt", ".md"):
        text = file_bytes.decode("utf-8", errors="ignore")

    else:
        raise ValueError("Supported formats: PDF, DOCX, TXT and MD.")

    text = text.replace("\x00", " ").strip()

    if len(text) < 50:
        raise ValueError(
            "The uploaded file does not contain enough readable text."
        )

    return text

# =========================================================
# CHUNKING
# =========================================================
def make_chunks(text, chunk_size=1100, overlap=180):
    text = " ".join(text.split())
    chunks = []
    start = 0

    while start < len(text):
        end = min(len(text), start + chunk_size)
        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = end - overlap

    return chunks

# =========================================================
# BUILD RESUME + JD RAG
# =========================================================
def rebuild_database():
    if not st.session_state.resume_text or not st.session_state.jd_text:
        st.session_state.db_ready = False
        return 0

    collection = get_chroma()

    existing = collection.get()
    if existing.get("ids"):
        collection.delete(ids=existing["ids"])

    docs = []
    metas = []
    ids = []

    documents = [
        ("resume", st.session_state.resume_text, st.session_state.resume_name),
        ("jd", st.session_state.jd_text, st.session_state.jd_name),
    ]

    for source_type, text, filename in documents:
        chunks = make_chunks(text)

        for index, chunk in enumerate(chunks):
            digest = hashlib.sha1(chunk.encode("utf-8")).hexdigest()[:12]
            ids.append(f"{source_type}_{index}_{digest}")
            docs.append(chunk)
            metas.append({
                "source_type": source_type,
                "source": filename or source_type,
                "chunk": index,
            })

    embeddings = get_embedder().encode(
        docs,
        normalize_embeddings=True,
        show_progress_bar=False,
    ).tolist()

    collection.add(
        ids=ids,
        documents=docs,
        embeddings=embeddings,
        metadatas=metas,
    )

    st.session_state.db_ready = True
    return len(docs)

# =========================================================
# RAG RETRIEVAL
# =========================================================
def retrieve(question, source_type=None, top_k=8):
    collection = get_chroma()

    if collection.count() == 0:
        return []

    query_vector = get_embedder().encode(
        [question],
        normalize_embeddings=True,
        show_progress_bar=False,
    )[0].tolist()

    kwargs = {
        "query_embeddings": [query_vector],
        "n_results": min(top_k, collection.count()),
        "include": ["documents", "metadatas", "distances"],
    }

    if source_type:
        kwargs["where"] = {"source_type": source_type}

    result = collection.query(**kwargs)

    documents = result.get("documents", [[]])[0]
    metadatas = result.get("metadatas", [[]])[0]
    distances = result.get("distances", [[]])[0]

    output = []

    for document, metadata, distance in zip(
        documents, metadatas, distances
    ):
        similarity = max(0.0, min(1.0, 1.0 - float(distance)))

        output.append({
            "text": document,
            "meta": metadata or {},
            "similarity": similarity,
        })

    return output

# =========================================================
# GEMINI
# =========================================================
def is_retryable(error):
    code = getattr(error, "code", None)
    text = str(error).lower()

    return (
        code in (429, 500, 502, 503, 504)
        or any(
            value in text
            for value in [
                "503",
                "429",
                "unavailable",
                "overloaded",
                "resource exhausted",
            ]
        )
    )

def call_gemini(prompt):
    client = get_gemini()
    last_error = None

    models = [
        (PRIMARY_MODEL, 2),
        (FALLBACK_MODEL, 1),
    ]

    for model, attempts in models:
        for attempt in range(attempts):
            try:
                response = client.models.generate_content(
                    model=model,
                    contents=prompt,
                )

                text = (getattr(response, "text", None) or "").strip()

                if text:
                    return text, model

                last_error = RuntimeError("Gemini returned an empty response.")

            except Exception as error:
                last_error = error

                if not is_retryable(error):
                    break

                if attempt < attempts - 1:
                    time.sleep(1.5 * (attempt + 1))

    raise last_error or RuntimeError("Gemini request failed.")

# =========================================================
# REAL RESUME-vs-JD SEMANTIC SCORE
# =========================================================
def semantic_match_score():
    resume_chunks = make_chunks(st.session_state.resume_text)
    jd_chunks = make_chunks(st.session_state.jd_text)

    if not resume_chunks or not jd_chunks:
        return 0.0

    embedder = get_embedder()

    resume_vectors = embedder.encode(
        resume_chunks,
        normalize_embeddings=True,
        show_progress_bar=False,
    )

    jd_vectors = embedder.encode(
        jd_chunks,
        normalize_embeddings=True,
        show_progress_bar=False,
    )

    # Cosine similarity because vectors are normalized.
    similarities = jd_vectors @ resume_vectors.T

    # For each JD chunk, find the closest resume chunk.
    best_matches = similarities.max(axis=1)

    # Convert semantic similarity into a coverage score.
    # This is deterministic and based only on the uploaded Resume/JD.
    coverage = ((best_matches - 0.25) / 0.65).clip(0, 1)

    return round(float(coverage.mean() * 100), 1)

# =========================================================
# LLM ANALYSIS
# =========================================================
def build_analysis():
    if not st.session_state.resume_text or not st.session_state.jd_text:
        return None

    score = semantic_match_score()

    resume_context = retrieve(
        "skills technologies education experience projects achievements",
        "resume",
        10,
    )

    jd_context = retrieve(
        "responsibilities required skills qualifications technologies experience",
        "jd",
        10,
    )

    resume_text = "\n\n".join(
        item["text"] for item in resume_context
    ) or st.session_state.resume_text[:10000]

    jd_text = "\n\n".join(
        item["text"] for item in jd_context
    ) or st.session_state.jd_text[:10000]

    prompt = f"""
You are CareerAI, a professional Resume-vs-Job-Description career analyst.

Analyze ONLY the supplied Resume and Job Description.

A deterministic semantic coverage score calculated from the actual
Resume/JD embeddings is:

{score}%

Use that exact number as the displayed Match Score.
Do not invent or change the score.

Return a useful report with these sections:

## 🎯 Match Score
Explain briefly what the score means.

## 💪 Strong Matches
List skills, technologies, experience, education or projects from
the resume that genuinely support the JD.

## ⚠️ Weak / Missing Areas
List important JD requirements that are absent or weakly supported
by the resume.

## ⭐ Resume Strengths for This JD
Explain the candidate's strongest evidence for this particular role.

## 🛠️ Improvement Priorities
Give practical actions the candidate can take.

## 🚀 Suggested Projects to Close the Gaps
Suggest realistic new projects specifically based on missing JD skills.
Clearly label these as SUGGESTED PROJECTS, not projects already completed.

## 🎯 Overall Recommendation
Give a concise assessment of fit and what the candidate should do next.

Rules:
- Never invent candidate skills, projects, education, employers or experience.
- A suggested project must never be presented as an existing project.
- Base strong/missing areas on the actual Resume and JD.
- Do not promise hiring or selection.
- Keep the report professional and useful for placement preparation.

RESUME:
{resume_text}

JOB DESCRIPTION:
{jd_text}
"""

    report, model = call_gemini(prompt)

    return {
        "score": score,
        "report": report,
        "model": model,
    }

# =========================================================
# CHAT RAG
# =========================================================
def answer_question(question):
    if not st.session_state.db_ready:
        return (
            "Please upload both your Resume and Job Description and "
            "click **Analyze Resume vs JD** first.",
            None,
        )

    hits = retrieve(question, None, 8)

    if not hits:
        return (
            "I could not find enough relevant information in the uploaded "
            "Resume and Job Description.",
            None,
        )

    context = "\n\n---\n\n".join(
        f"[{item['meta'].get('source_type', 'DOCUMENT').upper()}]\n"
        f"{item['text']}"
        for item in hits
    )

    prompt = f"""
You are CareerAI.

Answer the user's question using the supplied Resume and Job Description
context.

Rules:
- Ground factual claims in the supplied documents.
- Clearly distinguish candidate facts from JD requirements.
- Never invent a candidate skill, project, degree, employer or experience.
- If something is not supported by the uploaded documents, say so.
- For suitability questions, compare Resume evidence against JD requirements.
- For improvement/project questions, provide practical suggestions based
  on the actual gaps.
- Suggestions may use general career reasoning but must be labelled as
  suggestions.
- Do not expose chunks, embeddings, vector databases or internal RAG
  implementation details unless the user specifically asks.
- Answer directly and clearly.

RETRIEVED CONTEXT:
{context}

USER QUESTION:
{question}
"""

    try:
        answer, model = call_gemini(prompt)
        return answer, model
    except Exception:
        return (
            "Gemini is temporarily unavailable. Please try the question again.",
            None,
        )

# =========================================================
# SIDEBAR
# =========================================================
with st.sidebar:
    st.markdown("## 🎯 CareerAI")
    st.caption(f"Welcome, {st.session_state.user_name}")

    if st.button("🚪 Change User", use_container_width=True):
        for key in [
            "user_name",
            "messages",
            "resume_text",
            "resume_name",
            "resume_hash",
            "jd_text",
            "jd_name",
            "jd_hash",
            "analysis",
            "db_ready",
            "chroma_client",
            "chroma_collection",
        ]:
            if key in st.session_state:
                if key == "messages":
                    st.session_state[key] = []
                elif key in ("chroma_client", "chroma_collection"):
                    st.session_state[key] = None
                else:
                    st.session_state[key] = None
        st.rerun()

    st.divider()

    st.markdown("### 1️⃣ Upload Resume")
    st.markdown(
        '<div class="small-note">PDF, DOCX, TXT or MD</div>',
        unsafe_allow_html=True,
    )

    resume_file = st.file_uploader(
        "Resume",
        type=["pdf", "docx", "txt", "md"],
        key="resume_uploader",
        label_visibility="collapsed",
    )

    if resume_file is not None:
        raw = resume_file.getvalue()
        file_hash = hashlib.sha256(raw).hexdigest()

        if file_hash != st.session_state.resume_hash:
            try:
                with st.spinner("📄 Reading resume..."):
                    st.session_state.resume_text = extract_text(
                        resume_file.name,
                        raw,
                    )

                st.session_state.resume_name = resume_file.name
                st.session_state.resume_hash = file_hash
                st.session_state.analysis = None
                st.session_state.messages = []
                st.session_state.db_ready = False

                st.rerun()

            except Exception as error:
                st.error(f"Resume error: {error}")

    if st.session_state.resume_name:
        st.success(f"✓ {st.session_state.resume_name}")

    st.divider()

    st.markdown("### 2️⃣ Upload Job Description")
    st.markdown(
        '<div class="small-note">PDF, DOCX, TXT or MD</div>',
        unsafe_allow_html=True,
    )

    jd_file = st.file_uploader(
        "Job Description",
        type=["pdf", "docx", "txt", "md"],
        key="jd_uploader",
        label_visibility="collapsed",
    )

    if jd_file is not None:
        raw = jd_file.getvalue()
        file_hash = hashlib.sha256(raw).hexdigest()

        if file_hash != st.session_state.jd_hash:
            try:
                with st.spinner("💼 Reading job description..."):
                    st.session_state.jd_text = extract_text(
                        jd_file.name,
                        raw,
                    )

                st.session_state.jd_name = jd_file.name
                st.session_state.jd_hash = file_hash
                st.session_state.analysis = None
                st.session_state.messages = []
                st.session_state.db_ready = False

                st.rerun()

            except Exception as error:
                st.error(f"JD error: {error}")

    if st.session_state.jd_name:
        st.success(f"✓ {st.session_state.jd_name}")

    if st.session_state.resume_text and st.session_state.jd_text:
        st.divider()
        st.markdown("### 3️⃣ Analyze")

        if st.button(
            "🔍 Analyze Resume vs JD",
            use_container_width=True,
            type="primary",
        ):
            try:
                with st.status(
                    "Building Resume + JD RAG and generating analysis...",
                    expanded=False,
                ):
                    chunk_count = rebuild_database()
                    st.session_state.analysis = build_analysis()

                st.success(f"RAG ready • {chunk_count} knowledge sections")
                st.rerun()

            except Exception as error:
                st.error(f"Analysis error: {error}")

    if st.session_state.db_ready:
        st.success(
            f"✓ RAG ready: {get_chroma().count()} Resume + JD sections"
        )

    st.divider()

    if st.button(
        "🗑️ Clear Conversation",
        use_container_width=True,
    ):
        st.session_state.messages = []
        st.rerun()

# =========================================================
# MAIN HERO
# =========================================================
st.markdown(
    f"""
    <div class="hero">
        <div class="main-title">🎯 CareerAI</div>
        <div class="subtitle">RAG-Powered Resume & Job Description Intelligence</div>
        <div style="color:#f1f5f9;font-size:16px;line-height:1.7;">
            Welcome, <b>{st.session_state.user_name}</b>.
            Upload your Resume and Job Description to get a grounded
            match analysis and then ask CareerAI questions about the role.
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

# =========================================================
# BEFORE ANALYSIS
# =========================================================
if not st.session_state.resume_text or not st.session_state.jd_text:
    st.markdown(
        '<div class="section-title">📄 Start Your Career Analysis</div>',
        unsafe_allow_html=True,
    )

    c1, c2 = st.columns(2)

    with c1:
        st.markdown(
            """
            <div class="card">
                <h3>📄 Resume</h3>
                <p>Upload your CV in PDF, DOCX, TXT or MD format.</p>
                <p class="small-note">
                    CareerAI extracts and semantically indexes your
                    skills, education, experience and projects.
                </p>
            </div>
            """,
            unsafe_allow_html=True,
        )

    with c2:
        st.markdown(
            """
            <div class="card">
                <h3>💼 Job Description</h3>
                <p>Upload the actual JD for the position you want.</p>
                <p class="small-note">
                    No fixed job database is required. Your uploaded JD
                    becomes the second RAG knowledge source.
                </p>
            </div>
            """,
            unsafe_allow_html=True,
        )

# =========================================================
# ANALYSIS
# =========================================================
if st.session_state.resume_text and st.session_state.jd_text:
    st.markdown(
        '<div class="section-title">📊 Resume vs JD Analysis</div>',
        unsafe_allow_html=True,
    )

    if st.session_state.analysis:
        score = float(st.session_state.analysis["score"])
        score = max(0.0, min(100.0, score))

        st.markdown(
            f"""
            <div class="score-card">
                <div style="font-size:18px;color:#e2e8f0;font-weight:800;">
                    🎯 Resume–JD Match Score
                </div>
                <div class="big-score">{score:.1f}%</div>
                <div class="score-bar">
                    <div class="score-fill" style="width:{score:.1f}%"></div>
                </div>
                <div style="color:#cbd5e1;font-size:13px;">
                    Deterministic semantic coverage calculated from the
                    actual uploaded Resume and JD embeddings.
                    This is not a hiring probability.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.markdown(st.session_state.analysis["report"])

        st.caption(
            f"Analysis generated with {st.session_state.analysis['model']}. "
            "The displayed score is calculated from Resume/JD embeddings."
        )

    else:
        st.info(
            "Both documents are ready. Click **Analyze Resume vs JD** "
            "in the sidebar to build the RAG and generate your analysis."
        )

# =========================================================
# CHAT
# =========================================================
st.markdown(
    '<div class="section-title">💬 Ask CareerAI</div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="small-note">'
    'Ask: "What skills am I missing?", '
    '"Why is my match score low?", '
    '"Which projects should I build?", '
    '"Am I suitable for this role?", or any question supported by your Resume + JD.'
    '</div>',
    unsafe_allow_html=True,
)

for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

question = st.chat_input(
    "Ask CareerAI about your Resume + JD..."
)

if question:
    st.session_state.messages.append({
        "role": "user",
        "content": question,
    })

    with st.chat_message("user"):
        st.markdown(question)

    with st.chat_message("assistant"):
        with st.spinner("🧠 CareerAI is thinking..."):
            answer, model = answer_question(question)

        st.markdown(answer)

        if model:
            st.caption(f"AI response • {model}")

    st.session_state.messages.append({
        "role": "assistant",
        "content": answer,
        "model": model,
    })
